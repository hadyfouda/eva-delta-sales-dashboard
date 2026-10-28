import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from data import PRODUCTS, REPS_DATA

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Styling
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

unit_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
val_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
pts_fill = PatternFill(start_color="14532D", end_color="14532D", fill_type="solid")

total_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
total_font = Font(name="Calibri", size=11, bold=True, color="000000")

bold_font = Font(name="Calibri", size=11, bold=True)
regular_font = Font(name="Calibri", size=11)

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

# 1. Master Reps Sheet
ws_reps = wb.create_sheet(title="Reps_List")
ws_reps.append(["Rep Key", "Rep Name", "Start Quarter", "Status"])
for cell in ws_reps[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for r in REPS_DATA:
    status = "Active" if not r.get("noData") else "Pending Data"
    start_q = f"Q{r.get('startQ', 1)}"
    ws_reps.append([r["key"], r["name"], start_q, status])

# 2. Master Sales Sheet (Database format easy for Google Apps Script to query)
ws_sales = wb.create_sheet(title="Sales_Database")
headers = [
    "Rep_Key", "Rep_Name", "Quarter",
    "Product_Id", "Product_Code", "Product_Name",
    "Sales_Units", "Target_Units", "Unit_Deficit",
    "Sales_Value_EGP", "Target_Value_EGP", "Value_Deficit_EGP",
    "Sales_Points", "Target_Points", "Points_Deficit",
    "Ratio_Pct"
]
ws_sales.append(headers)
for cell in ws_sales[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for rep in REPS_DATA:
    if rep.get("noData"):
        continue
    for q in ["Q1", "Q2", "Q3"]:
        qd = rep["quarters"].get(q)
        if not qd or not qd.get("products"):
            continue
        for p in qd["products"]:
            prod = next((pr for pr in PRODUCTS if pr["id"] == p["pid"]), None)
            su = p.get("salesUnits")
            tu = p.get("targetUnits")
            u_def = (tu - su) if (tu is not None and su is not None) else None

            sv = p.get("salesValue")
            tv = p.get("targetValue")
            v_def = (tv - sv) if (tv is not None and sv is not None) else None

            sp = p.get("salesPoints")
            tp = p.get("targetPoints")
            p_def = (tp - sp) if (tp is not None and sp is not None) else None

            ratio = p.get("ratio")

            ws_sales.append([
                rep["key"], rep["name"], q,
                p["pid"], prod["code"] if prod else "", prod["name"] if prod else "",
                su, tu, u_def,
                sv, tv, v_def,
                sp, tp, p_def,
                ratio
            ])

# Auto-adjust column widths
for ws in [ws_reps, ws_sales]:
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

out_xlsx = r"C:\Users\ALMadina\.gemini\antigravity\scratch\delta-team-dashboard\Delta_Team_Database.xlsx"
wb.save(out_xlsx)
print("Saved Master Excel Database to:", out_xlsx)
